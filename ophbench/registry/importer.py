from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

import yaml
from openpyxl import load_workbook

from .schemas import CheckpointRecord, ModelRecord

MODEL_ID_MAP = {
    "RETFound": "retfound",
    "VisionFM": "visionfm",
    "EyeCLIP": "eyeclip",
    "FLAIR": "flair",
    "RetiZero": "retizero",
    "MIRAGE": "mirage",
    "RETFound-Green": "retfound-green",
    "VisionUnite": "visionunite",
    "FMUE": "fmue",
    "KeepFIT": "keepfit",
    "RET-CLIP": "ret-clip",
    "ViLReF": "vilref",
    "PRETI": "preti",
    "UrFound": "urfound",
    "DERETFound": "deretfound",
}
EXPECTED_CHECKPOINT_IDS = [
    "retfound-cfp",
    "retfound-oct",
    "visionfm-fundus",
    "visionfm-oct",
    "visionfm-ffa",
    "visionfm-ultrasound",
    "visionfm-external-eye",
    "visionfm-slit-lamp",
    "visionfm-mri",
    "visionfm-ubm",
    "eyeclip-default",
    "flair-default",
    "retizero-default",
    "mirage-base",
    "mirage-large",
    "retfound-green-v0.1",
    "visionunite-default",
    "fmue-default",
    "keepfit-flair-mmretinal-cfp",
    "keepfit-half-flair-mmretinal-cfp",
    "keepfit-ffa-ir-mmretinal-ffa",
    "ret-clip-default",
    "vilref-default",
    "preti-default",
    "urfound-default",
    "deretfound-pretraining",
    "deretfound-sd-retina",
]
PHASES = {
    **{
        x: "phase1_image_encoder"
        for x in ["retfound", "visionfm", "mirage", "retfound-green", "preti", "deretfound"]
    },
    "fmue": "phase1_specialized",
    **{
        x: "phase2_vision_language"
        for x in [
            "eyeclip",
            "flair",
            "retizero",
            "visionunite",
            "keepfit",
            "ret-clip",
            "vilref",
            "urfound",
        ]
    },
}
MODALITY_ORDER = [
    "CFP",
    "OCT",
    "OCTA",
    "FFA",
    "SLO",
    "ICGA",
    "FAF",
    "ultrasound",
    "external_eye",
    "slit_lamp",
    "specular_microscopy",
    "MRI",
    "UBM",
    "CT",
    "RetCam",
    "PET",
    "X_ray",
    "text",
    "retinal_layer_pseudolabels",
    "other",
]


@dataclass(frozen=True)
class ImportResult:
    model_count: int
    checkpoint_count: int


def _clean(value):
    return " ".join(str(value or "").split()).strip()


def _modalities(text):
    return _parse_modalities(text)[0]


def _parse_modalities(text):
    raw = _clean(text)
    aliases = [
        ("CFP", "CFP"),
        ("OCTA", "OCTA"),
        ("OCT", "OCT"),
        ("FFA", "FFA"),
        ("FA", "FFA"),
        ("SLO", "SLO"),
        ("SLP", "slit_lamp"),
        ("ICGA", "ICGA"),
        ("FAF", "FAF"),
        ("B超", "ultrasound"),
        ("Ultrasound", "ultrasound"),
        ("OUS", "ultrasound"),
        ("外眼", "external_eye"),
        ("External Eye", "external_eye"),
        ("EEP", "external_eye"),
        ("裂隙灯", "slit_lamp"),
        ("Slit Lamp", "slit_lamp"),
        ("SM", "specular_microscopy"),
        ("MRI", "MRI"),
        ("UBM", "UBM"),
        ("CT", "CT"),
        ("RetCam", "RetCam"),
        ("PET", "PET"),
        ("X-ray", "X_ray"),
        ("文本", "text"),
        ("报告", "text"),
        ("Pseudo-label", "retinal_layer_pseudolabels"),
    ]
    values = []
    for needle, canonical in aliases:
        if needle.isascii() and needle.isupper():
            matched = re.search(
                rf"(?<![A-Za-z]){re.escape(needle)}(?![A-Za-z])", raw, re.IGNORECASE
            )
        else:
            matched = needle.lower() in raw.lower()
        if matched and canonical not in values:
            values.append(canonical)
    known_tokens = {needle.upper() for needle, _ in aliases if needle.isascii()}
    explicit_tokens = set(re.findall(r"(?<![A-Za-z])[A-Z][A-Z0-9-]{1,}(?![A-Za-z])", raw))
    unmapped = sorted(token for token in explicit_tokens if token.upper() not in known_tokens)
    if unmapped or not values:
        values.append("other")
    values.sort(key=MODALITY_ORDER.index)
    return values, unmapped


def _capabilities(text):
    rules = [
        ("特征", "feature_extraction"),
        ("分类", "classification"),
        ("风险", "risk_prediction"),
        ("预后", "prognosis"),
        ("分割", "segmentation"),
        ("检测", "detection"),
        ("关键点", "keypoint_localization"),
        ("零样本", "zero_shot_classification"),
        ("少样本", "few_shot_classification"),
        ("检索", "retrieval"),
        ("对齐", "multimodal_alignment"),
        ("问答", "vqa"),
        ("VQA", "vqa"),
        ("报告生成", "report_generation"),
        ("不确定", "uncertainty_estimation"),
        ("开放集", "ood_detection"),
    ]
    result = []
    for needle, value in rules:
        if needle.lower() in text.lower() and value not in result:
            result.append(value)
    return result


def _publication(value):
    text = _clean(value)
    match = re.match(r"(\d{4})[，,]\s*(.+)", text)
    year, venue = (int(match.group(1)), match.group(2)) if match else (None, text)
    lower = venue.lower()
    ptype = (
        "conference"
        if "miccai" in lower
        else "preprint"
        if "arxiv" in lower
        else "journal"
        if venue
        else "unknown"
    )
    return year, ptype, venue


def _provider(url):
    if not url:
        return "unknown"
    if "huggingface.co" in url:
        return "huggingface"
    if "drive.google.com" in url:
        return "google_drive"
    if "zenodo.org" in url:
        return "zenodo"
    if "github.com" in url and "/releases/" in url:
        return "github_release"
    return "other"


def _yaml(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")


def import_seed(input_path: Path, output_root: Path = Path("."), imported_at: str | None = None):
    imported_at = imported_at or (
        datetime.now(timezone.utc)
        .replace(microsecond=0)
        .isoformat()
        .replace("+00:00", "Z")
    )
    workbook = load_workbook(input_path, data_only=True, read_only=True)
    main = workbook["眼科大模型"]
    headers = [_clean(c.value) for c in next(main.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row_number, cells in enumerate(main.iter_rows(min_row=2, values_only=True), start=2):
        row = dict(zip(headers, cells, strict=True))
        rows.append((row_number, row))
        name = _clean(row["模型"])
        model_id = MODEL_ID_MAP[name]
        year, publication_type, venue = _publication(row["年份/文献状态"])
        modality_text = _clean(row["预训练模态"])
        modalities, unmapped_modalities = _parse_modalities(modality_text)
        model = {
            "schema_version": "1.0",
            "model_id": model_id,
            "model_name": name,
            "year": year,
            "publication_type": publication_type,
            "venue": venue,
            "model_category": _clean(row["模型类别"]),
            "modalities": modalities,
            "architecture": _clean(row["核心架构"]),
            "pretraining_data_summary": _clean(row["预训练数据规模/来源"]),
            "pretraining_strategy": _clean(row["预训练策略"]),
            "reported_summary": _clean(row["主要表现/特点"]),
            "paper_url": _clean(row["论文链接"]) or None,
            "code_url": _clean(row["GitHub仓库"]) or None,
            "code_available": _clean(row["是否提供代码"]) == "是",
            "capabilities": _capabilities(_clean(row["适用任务"])),
            "benchmark_tracks": [],
            "runtime_phase": PHASES[model_id],
            "license": None,
            "license_verified": False,
            "verification_status": "seed_unverified",
            "implementation": {
                "adapter_status": "not_started",
                "smoke_test_status": "not_run",
                "benchmark_status": "not_run",
            },
            "reported_tasks_text": _clean(row["适用任务"]),
            "provenance": {
                "source_file": input_path.name,
                "source_sheet": "眼科大模型",
                "source_row": row_number,
                "imported_at": imported_at,
            },
            "notes": [f"Original modality text: {modality_text}"],
        }
        if unmapped_modalities:
            model["notes"].append(
                f"Unmapped modality tokens: {', '.join(unmapped_modalities)}"
            )
        _yaml(output_root / "registry/models" / f"{model_id}.yaml", model)

    multi = workbook["有多个权重的模型"]
    multi_rows = {
        r: [_clean(c) for c in values]
        for r, values in enumerate(multi.iter_rows(values_only=True), start=1)
    }
    specs = [
        ("retfound-cfp", "retfound", "CFP", ["CFP"], multi_rows[1][1], 1),
        ("retfound-oct", "retfound", "OCT", ["OCT"], multi_rows[1][2], 1),
        *[
            (cid, "visionfm", name, mods, multi_rows[2][i], 2)
            for i, (cid, name, mods) in enumerate(
                [
                    ("visionfm-fundus", "Fundus", ["CFP"]),
                    ("visionfm-oct", "OCT", ["OCT"]),
                    ("visionfm-ffa", "FFA", ["FFA"]),
                    ("visionfm-ultrasound", "Ultrasound", ["ultrasound"]),
                    ("visionfm-external-eye", "External Eye", ["external_eye"]),
                    ("visionfm-slit-lamp", "Slit Lamp", ["slit_lamp"]),
                    ("visionfm-mri", "MRI", ["MRI"]),
                    ("visionfm-ubm", "UBM", ["UBM"]),
                ],
                start=1,
            )
        ],
        ("mirage-base", "mirage", "Base", ["OCT", "SLO"], multi_rows[3][1], 3),
        ("mirage-large", "mirage", "Large", ["OCT", "SLO"], multi_rows[3][2], 3),
        (
            "keepfit-flair-mmretinal-cfp",
            "keepfit",
            "FLAIR + MM-Retinal CFP",
            ["CFP"],
            multi_rows[4][1],
            4,
        ),
        (
            "keepfit-half-flair-mmretinal-cfp",
            "keepfit",
            "Half FLAIR + MM-Retinal CFP",
            ["CFP"],
            multi_rows[4][2],
            4,
        ),
        (
            "keepfit-ffa-ir-mmretinal-ffa",
            "keepfit",
            "FFA-IR + MM-Retinal FFA",
            ["FFA"],
            multi_rows[4][3],
            4,
        ),
        ("deretfound-pretraining", "deretfound", "Pretraining", ["CFP"], multi_rows[5][1], 5),
        (
            "deretfound-sd-retina",
            "deretfound",
            "Stable Diffusion Retina",
            ["CFP"],
            multi_rows[5][2],
            5,
        ),
    ]
    multi_models = {s[1] for s in specs}
    for row_number, row in rows:
        model_id = MODEL_ID_MAP[_clean(row["模型"])]
        if model_id not in multi_models:
            name = "v0.1" if model_id == "retfound-green" else "Default"
            specs.append(
                (
                    f"{model_id}-v0.1" if model_id == "retfound-green" else f"{model_id}-default",
                    model_id,
                    name,
                    _modalities(_clean(row["预训练模态"])),
                    _clean(row["权重下载链接"]),
                    row_number,
                )
            )
    row_lookup = {MODEL_ID_MAP[_clean(row["模型"])]: (n, row) for n, row in rows}
    for checkpoint_id, model_id, name, modalities, url, source_row in specs:
        main_row_num, main_row = row_lookup[model_id]
        status_text = _clean(main_row["权重状态"])
        auth = "登录" in _clean(main_row["权重获取方式"]) or "申请" in status_text
        checkpoint = {
            "schema_version": "1.0",
            "checkpoint_id": checkpoint_id,
            "model_id": model_id,
            "checkpoint_name": name,
            "modalities": modalities,
            "weight_url": url or None,
            "provider": _provider(url),
            "access_type": "auth_required" if auth else "open" if url else "unknown",
            "requires_auth": auth,
            "redistribution_allowed": "unknown",
            "license": None,
            "framework": None,
            "input_size": None,
            "normalization": None,
            "embedding_dim": None,
            "sha256": None,
            "source_commit": None,
            "verification_status": "seed_unverified",
            "last_verified": None,
            "provenance": {
                "source_file": input_path.name,
                "source_sheet": "有多个权重的模型" if model_id in multi_models else "眼科大模型",
                "source_row": source_row if model_id in multi_models else main_row_num,
                "imported_at": imported_at,
            },
            "notes": ["Seed metadata only; URL and redistribution terms have not been verified."],
        }
        _yaml(output_root / "registry/checkpoints" / f"{checkpoint_id}.yaml", checkpoint)
    seed_dir = output_root / "seed"
    seed_dir.mkdir(parents=True, exist_ok=True)
    manifest = {
        "source_file": input_path.name,
        "sha256": hashlib.sha256(input_path.read_bytes()).hexdigest(),
        "imported_at": imported_at,
        "source_sheets": workbook.sheetnames,
        "expected_model_count": 15,
        "expected_checkpoint_count": 27,
        "importer_version": "0.1.0",
    }
    (seed_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    schema_dir = output_root / "registry/schemas"
    schema_dir.mkdir(parents=True, exist_ok=True)
    (schema_dir / "model.schema.json").write_text(
        json.dumps(ModelRecord.model_json_schema(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    (schema_dir / "checkpoint.schema.json").write_text(
        json.dumps(CheckpointRecord.model_json_schema(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return ImportResult(len(rows), len(specs))
